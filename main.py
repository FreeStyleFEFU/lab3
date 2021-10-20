import csv
from openpyxl import load_workbook


def unique_words():
    with open('files/input.txt', "r", encoding="utf8") as file:
        contents = file.readlines()
        unique_words = []
        for row in contents:
            words_in_row = row.split()
            for word in words_in_row:
                if word not in unique_words:
                    unique_words.append(word)
        print(f'Уникальных слов в тексте {len(unique_words)}')


def quantity_every_word():
    with open('files/input.txt', "r", encoding="utf8") as file:
        contents = file.readlines()
        unique_words = {}
        for row in contents:
            words_in_row = row.split()
            for word in words_in_row:
                if word in unique_words:
                    unique_words[word] += 1
                else:
                    unique_words[word] = 1

        quantity_every_word = dict((sorted(unique_words.items(), key=lambda x: x[1], reverse=True)))

        words_string = ''

        for word in quantity_every_word.keys():
            words_string += word + ' '

        print(words_string)


def longest_rows():
    with open('files/input (4).txt', "r", encoding="utf8") as file:
        contents = file.readlines()
        max_length = max(len(row) for row in contents)

        for row in contents:
            if len(row) == max_length:
                print(row)


def reverse_text():
    with open('files/input (6).txt', "r", encoding="utf8") as file:
        contents = file.readlines()
        contents.reverse()
        for row in contents:
            print(''.join(reversed(row)))


def sort_salary():
    with open('files/input.csv') as file:
        file_reader = csv.reader(file, delimiter=";")
        company_salary = {}
        for row in file_reader:
            company_salary[row[0]] = int(row[1])

        sorted_company_salary = dict((sorted(company_salary.items(), key=lambda x: x[1])))

        for company in sorted_company_salary:
            print(f'{company}: {sorted_company_salary[company]}')


def sort_calories():
    wb = load_workbook('files/trekking1.xlsx')
    print(wb.sheetnames)

sort_calories()